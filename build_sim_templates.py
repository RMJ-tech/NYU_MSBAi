import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook()
INP=PatternFill("solid",fgColor="DDEBF7")     # input (blue)
ASM=PatternFill("solid",fgColor="FFFF00")     # CB assumption (yellow)
FOR=PatternFill("solid",fgColor="FFC000")     # CB forecast (orange)
HDR=PatternFill("solid",fgColor="1F4E78"); GREY=PatternFill("solid",fgColor="F2F2F2")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=10)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)

def C(ws,ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    c=ws[ref]
    if val is not None: c.value=val
    if fill: c.fill=fill
    if font: c.font=font
    if align: c.alignment=align
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

# ============ Sheet 3A ============
ws=wb.active; ws.title="3A_MutualFunds"; ws.sheet_view.showGridLines=False
ws.merge_cells("A1:D1")
C(ws,"A1","3A. Mutual Funds — P(best of 100 funds beats market ≥ 11 of 13 yrs)",font=title_f,border=False,align=left)
ws.merge_cells("A2:D2")
C(ws,"A2","Blue = inputs.  Yellow = CB Assumption cells.  Orange = CB Forecast cell.  (=CB.* shows #NAME? until Crystal Ball is loaded.)",font=note,border=False,align=left)

C(ws,"A4","P(beat market / year)",font=bold,align=left); C(ws,"B4",0.5,fill=INP,align=ctr,fmt="0.00")
C(ws,"A5","Years",font=bold,align=left);                 C(ws,"B5",13,fill=INP,align=ctr)
C(ws,"A6","Number of funds",font=bold,align=left);       C(ws,"B6",100,fill=INP,align=ctr)
C(ws,"A7","Threshold (≥)",font=bold,align=left);         C(ws,"B7",11,fill=INP,align=ctr)

C(ws,"A9","Fund #",fill=HDR,font=hf,align=ctr)
C(ws,"B9","Years beaten  =CB.Binomial(p, years)",fill=HDR,font=hf,align=ctr)
for k in range(100):
    r=10+k
    C(ws,f"A{r}",k+1,align=ctr)
    C(ws,f"B{r}","=CB.Binomial($B$4,$B$5)",fill=ASM,align=ctr,font=mono)   # ASSUMPTION

C(ws,"A111","Best of 100 (max)",font=bold,align=left); C(ws,"B111","=MAX(B10:B109)",fill=GREY,font=bold,align=ctr)
C(ws,"A112","FORECAST → best ≥ 11 ? (1/0)",font=bold,align=left)
C(ws,"B112","=IF(B111>=$B$7,1,0)",fill=FOR,font=Font(bold=True,size=12),align=ctr)   # FORECAST

C(ws,"A114","HOW TO USE",font=Font(bold=True,color="1F4E78"),border=False,align=left)
for i,t in enumerate([
 "1. Define B10:B109 as Assumptions (typing =CB.Binomial auto-registers them).",
 "2. Define B112 as the Forecast cell.",
 "3. Run Preferences → 10,000 trials → Run.",
 "4. B112 forecast MEAN = the probability (P best ≥ 11).  Capture frequency chart + statistics table.",
 "Tip: verify CB.Binomial argument order (prob, trials) in the CB function wizard.",
]):
    C(ws,f"A{115+i}",t,border=False,align=left,font=note); ws.merge_cells(f"A{115+i}:D{115+i}")
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=34

# ============ Sheet 3B ============
ws2=wb.create_sheet("3B_MontyHall"); ws2.sheet_view.showGridLines=False
ws2.merge_cells("A1:D1")
C(ws2,"A1","3B. Let's Make a Deal — should you switch?",font=title_f,border=False,align=left)
ws2.merge_cells("A2:D2")
C(ws2,"A2","Yellow = CB Assumption.  Orange = CB Forecast.  Each trial = one play of the game.",font=note,border=False,align=left)

C(ws2,"A4","Prize door (1-3)",font=bold,align=left);  C(ws2,"B4","=CB.DiscreteUniform(1,3)",fill=ASM,align=ctr,font=mono)  # ASSUMPTION
C(ws2,"A5","Chosen door (always 1)",font=bold,align=left); C(ws2,"B5",1,fill=INP,align=ctr)

C(ws2,"A7","FORECAST → Win if NO switch (1/0)",font=bold,align=left)
C(ws2,"B7","=IF(B4=B5,1,0)",fill=FOR,font=Font(bold=True,size=12),align=ctr)         # FORECAST
C(ws2,"A8","FORECAST → Win if SWITCH (1/0)",font=bold,align=left)
C(ws2,"B8","=IF(B4<>B5,1,0)",fill=FOR,font=Font(bold=True,size=12),align=ctr)        # FORECAST

C(ws2,"A10","Wins, NO switch (× 10,000 plays)",font=bold,align=left); C(ws2,"B10","= mean(B7) × 10000",border=True,align=ctr,font=note)
C(ws2,"A11","Wins, SWITCH (× 5,000 plays)",font=bold,align=left);     C(ws2,"B11","= mean(B8) × 5000",border=True,align=ctr,font=note)

C(ws2,"A13","HOW TO USE",font=Font(bold=True,color="1F4E78"),border=False,align=left)
for i,t in enumerate([
 "1. Define B4 as Assumption (=CB.DiscreteUniform auto-registers).",
 "2. Define B7 AND B8 as Forecast cells.",
 "3. No-switch: Run 10,000 trials → wins ≈ B7 mean × 10,000.",
 "4. Switch: Run 5,000 trials → wins ≈ B8 mean × 5,000.  (mean is same regardless of trial count; only the count scales.)",
 "5. Compare the two win rates → decide whether to switch.  Capture frequency charts + statistics tables.",
]):
    C(ws2,f"A{14+i}",t,border=False,align=left,font=note); ws2.merge_cells(f"A{14+i}:D{14+i}")
ws2.column_dimensions["A"].width=34; ws2.column_dimensions["B"].width=30

out="/home/user/NYU_MSBAi/Simulation_Q3_Templates.xlsx"
wb.save(out); print("Wrote",out,"sheets:",wb.sheetnames)
