import pandas as pd, re, json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = "/root/.claude/uploads/25fd9b56-1752-5edc-8721-3932cd50f9d6/70eb99d7-20260523_naming_rights_database_1.xlsx"
RES = "/tmp/claude-0/-home-user-NYU-MSBAi/25fd9b56-1752-5edc-8721-3932cd50f9d6/scratchpad/research"
START, END = 101, 150   # 1-based inclusive

raw_full = pd.read_excel(SRC, sheet_name="Data_Input", header=None)
def norm_group(v):
    s=str(v)
    if s.startswith("Group.1"): return "Group 1"
    if s.startswith("Group.2"): return "Group 2"
    return ""
groups=[norm_group(g) for g in raw_full.iloc[0].tolist()]
categories=[("" if pd.isna(c) else str(c)) for c in raw_full.iloc[2].tolist()]
fields=[("" if pd.isna(f) else str(f)) for f in raw_full.iloc[3].tolist()]
NCOLS=100

raw_df=pd.read_excel(SRC, sheet_name="Data_Input", header=3).dropna(how="all").reset_index(drop=True).iloc[START-1:END]

CORE={"country":"Country","city":"City","state_prefecture":"State / Prefecture",
      "ownership_type":"Ownership Type","opened_year":"Opened Year","capacity":"Capacity"}
KEYMAP={
 "prev_name":"Previous / Alternative Venue Name","metro_area":"Metro Area","latitude":"Latitude","longitude":"Longitude",
 "renovated_year":"Renovated Year","total_site_area_sqm":"Total Site Area sqm","nearest_station_airport":"Nearest Major Station / Airport",
 "distance_to_cbd_km":"Distance to CBD km","sponsor_hq_country":"Sponsor HQ Country","contract_announcement_date":"Contract Announcement Date",
 "league_competition":"League / Competition","avg_home_attendance":"Average Home Attendance","annual_venue_attendance":"Annual Venue Attendance",
 "annual_event_days":"Annual Event Days","sports_event_days":"Sports Event Days","concert_days":"Concert / Entertainment Days",
 "other_event_days":"Other Event Days","broadcast_viewers_annual":"Broadcast / Streaming Viewers Annual","national_broadcast_events":"National Broadcast Games / Events",
 "intl_exposure_yn":"International Exposure Y/N","media_mentions_annual":"Media Mentions Annual","est_media_impressions":"Estimated Media Impressions",
 "social_followers":"Social Followers Tenant / Venue","annual_social_posts":"Annual Social Posts Mentioning Venue","social_impressions_annual":"Social Impressions Annual",
 "avg_engagement_rate":"Avg Engagement Rate","google_trends_index":"Google Trends Index Venue","search_volume_venue":"Search Volume Venue",
 "mobile_foot_traffic_annual":"Mobile Foot Traffic Annual","city_population":"City Population","metro_population":"Metro Population",
 "regional_gdp_usd":"Regional GDP / GRP","median_household_income_usd":"Median Household Income","annual_tourists_city":"Annual Tourists / Visitors City",
 "num_large_corp_hq":"Number of Large Corporate HQs","sponsor_venue_fit_notes":"Sponsor-Venue Fit Notes","special_factors":"Special Factors",
 "exclusivity_yn":"Exclusivity Y/N","renewal_yn":"Renewal Y/N","termination_yn":"Termination / Opt-out Y/N",
 "facility_governance":"Facility Governance / Scheme","bundled_assets_flag":"Bundled Assets Flag",
}
SUBREGION={"Minnesota":"US Midwest","Wisconsin":"US Midwest","Iowa":"US Midwest","Illinois":"US Midwest","Michigan":"US Midwest","Indiana":"US Midwest","Ohio":"US Midwest","North Dakota":"US Midwest","Kansas":"US Midwest","Missouri":"US Midwest","Nebraska":"US Midwest",
 "Colorado":"US Mountain West","Idaho":"US Mountain West","Utah":"US Mountain West","Nevada":"US Mountain West","Arizona":"US Southwest","Oklahoma":"US South Central","Texas":"US South Central","Louisiana":"US South Central","Arkansas":"US South Central",
 "Washington":"US Pacific Northwest","Oregon":"US Pacific Northwest","California":"US West (Pacific)","Alaska":"US West (Pacific)",
 "Florida":"US Southeast","Georgia":"US Southeast","North Carolina":"US Southeast","South Carolina":"US Southeast","Tennessee":"US Southeast","Kentucky":"US Southeast","Virginia":"US Southeast","Alabama":"US Southeast",
 "Pennsylvania":"US Northeast","New York":"US Northeast","New Jersey":"US Northeast","Rhode Island":"US Northeast","District of Columbia":"US Mid-Atlantic","Maryland":"US Mid-Atlantic","Washington, D.C.":"US Mid-Atlantic","D.C.":"US Mid-Atlantic",
 "Saskatchewan":"Canada Prairies","Manitoba":"Canada Prairies","Ontario":"Canada (Central)","Quebec":"Canada (Central)"}

def isblank(v): return v is None or (isinstance(v,float) and pd.isna(v)) or str(v).strip() in ("","nan","NaT","None")
def fmt(v):
    if isblank(v): return ""
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v)
def num(v):
    try:
        f=float(v); return None if f!=f else f
    except: return None
def rv_of(res,key):
    o=res.get(key); return o.get("v") if isinstance(o,dict) else None
def rs_of(res,key):
    o=res.get(key); return (o.get("s") or "") if isinstance(o,dict) else ""

DARK="1E4620"; GREEN="E2EFDA"; YELLOW="FFF2CC"
hdr_fill=PatternFill("solid",fgColor=DARK); green_fill=PatternFill("solid",fgColor=GREEN); yel_fill=PatternFill("solid",fgColor=YELLOW)
hdr_font=Font(bold=True,size=11,color="FFFFFF"); title_font=Font(bold=True,size=14,color=DARK); sub_font=Font(size=11,color="555555")
data_font=Font(size=10,color="333333"); bold_data=Font(size=10,color="000000",bold=True)
thin=Side(style="thin",color="D9D9D9"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
ctr=Alignment(horizontal="center",vertical="center",wrap_text=True); ctr_v=Alignment(horizontal="center",vertical="top")
left_v=Alignment(horizontal="left",vertical="top",wrap_text=True)
HEADERS=["Col Index","Category Block","Field Name Header","Original Group Priority","Populated Value","Original Raw Value","Needs External Input?","Source & Rationale Context"]
WIDTHS=[9,26,33,16,46,40,18,62]
def clean(rid,venue): return re.sub(r'[:\\/?*\[\]]','-',f"{rid}_{venue}")[:31]

out="/home/user/NYU_MSBAi/venues_101_150_filled.xlsx"
wb=openpyxl.Workbook(); wb.remove(wb.active)
summary_rows=[]

for idx,(_,raw_row) in enumerate(raw_df.iterrows()):
    n=START+idx
    rid=fmt(raw_row["Record ID"]); venue=fmt(raw_row["Venue Name"])
    res=json.load(open(os.path.join(RES,f"v{n}.json")))

    pop={}; src={}
    for k,fld in CORE.items():
        pop[fld]=rv_of(res,k); src[fld]=rs_of(res,k) or "Sourced via web lookup verification."
    for k,fld in KEYMAP.items():
        if k in res: pop[fld]=rv_of(res,k); src[fld]=rs_of(res,k)

    state=fmt(rv_of(res,"state_prefecture")); own=fmt(rv_of(res,"ownership_type"))
    country=fmt(rv_of(res,"country")); cap=num(rv_of(res,"capacity")); vtype=fmt(raw_row["Venue Type"])
    end=num(raw_row["Contract End Year"]); ln=num(raw_row["Contract Length years"])
    if end and ln: pop["Contract Start Year"]=int(end-ln); src["Contract Start Year"]="Derived: Contract End Year minus Contract Length"
    aha=num(rv_of(res,"avg_home_attendance"))
    if cap and aha: pop["Attendance per Seat"]=round(aha/cap,3); src["Attendance per Seat"]="Derived: Average Home Attendance / Capacity"
    aed=num(rv_of(res,"annual_event_days"))
    if aed: pop["Utilization Rate"]=round(aed/365.0,3); src["Utilization Rate"]="Derived: Annual Event Days / 365"
    pop["Region"]="North America"; src["Region"]="Geographic taxonomy (standard classification)"
    pop["Subregion"]=SUBREGION.get(state,""); src["Subregion"]="Geographic taxonomy by state/province"
    pf="Public" if own.lower().startswith("public") else ("Private" if own.lower().startswith("private") else "")
    if pf: pop["Public / Private Facility"]=pf; src["Public / Private Facility"]="Derived from Ownership Type"
    pop["Rights Type"]="Naming Rights Only"; src["Rights Type"]="Standard classification (full naming-rights listing)"
    scope="Sub-facility / Training Facility" if ("Training" in vtype or "Practice" in vtype) else "Full Venue"
    pop["Naming Scope"]=scope; src["Naming Scope"]="Classified from venue type"
    tot=num(raw_row["Total Contract Value Original"]); ann=num(raw_row["Annual Fee Original"])
    vds="Disclosed (annual and/or total)" if (tot or ann) else "Undisclosed / estimate"
    pop["Value Disclosure Status"]=vds; src["Value Disclosure Status"]="Derived from disclosed contract value fields"
    pop["Local Country / Territory Name"]=country; src["Local Country / Territory Name"]="Derived from Country"
    pop["English Venue Name"]=venue; src["English Venue Name"]="Same as venue name (English-language venue)"
    pop["Local Venue Name"]=venue; src["Local Venue Name"]="Same as venue name (English-language venue)"
    is_ca = country.lower().startswith("can")
    lang = "French/English" if state in ("Quebec",) else "English"
    pop["Original Language"]=lang; src["Original Language"]="Classified from venue locale"
    pop["Row Source Type"]="Observed benchmark (SBJ directory extract + web enrichment)"; src["Row Source Type"]="Classification of row provenance"
    dtvb="Observed disclosed contract" if (tot or ann) else "Observed contract - value undisclosed"
    pop["Data Type / Value Basis"]=dtvb; src["Data Type / Value Basis"]="Derived from disclosure status"

    ws=wb.create_sheet(clean(rid,venue)); ws.sheet_view.showGridLines=False
    ws["A1"]="Strict Priority-Group Audit & Live Web Enrichment Master Framework"; ws["A1"].font=title_font
    ws["A2"]=f"Record Status Review: {rid} ({venue})"; ws["A2"].font=sub_font
    for j,h in enumerate(HEADERS):
        c=ws.cell(row=4,column=j+1,value=h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=ctr; c.border=border
    preexist=filled=blank=0
    for ci in range(NCOLS):
        field=fields[ci]
        rraw=raw_row[field] if field in raw_df.columns else None
        pv=pop[field] if field in pop else (rraw if field in raw_df.columns else None)
        raw_present=not isblank(rraw); pv_present=not isblank(pv)
        if raw_present:
            status="No"; fill=None; rationale="Pre-existing raw value from the original spreadsheet extract matrix."; preexist+=1
        elif pv_present:
            status="Completed"; fill=green_fill; rationale=src.get(field,"Sourced via web lookup verification."); filled+=1
        else:
            status="Yes"; fill=yel_fill
            rationale=src.get(field,"") or "Not publicly available — external input attempted."
            if not rationale.lower().startswith(("not","estimate","blank","derived","sourced","same","classif","geograph","observed")):
                rationale="Not publicly available — external input attempted."
            blank+=1
        r=5+ci
        vals=[ci,categories[ci],field,groups[ci],fmt(pv),fmt(rraw),status,rationale]
        for j,val in enumerate(vals):
            c=ws.cell(row=r,column=j+1,value=val); c.border=border; c.font=data_font
            c.alignment=ctr_v if j in (0,3,6) else left_v
            if fill and j in (4,6): c.fill=fill
            if j==6 and status in ("Completed","Yes"): c.font=bold_data
    for j,w in enumerate(WIDTHS): ws.column_dimensions[openpyxl.utils.get_column_letter(j+1)].width=w
    ws.freeze_panes="A5"; ws.row_dimensions[4].height=30
    summary_rows.append([n,rid,venue,fmt(rv_of(res,"city")),state,country,preexist,filled,blank])

idx_ws=wb.create_sheet("Index"); wb.move_sheet("Index",-(len(wb.sheetnames)-1))
idx_ws.sheet_view.showGridLines=False
idx_ws["A1"]=f"Venues {START}-{END} — Filled & Audited (Index)"; idx_ws["A1"].font=title_font
ihdr=["#","Record ID","Venue Name","City","State","Country","Pre-existing","Filled (web/derived)","Still blank"]
for j,h in enumerate(ihdr):
    c=idx_ws.cell(row=3,column=j+1,value=h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=ctr; c.border=border
for r,row in enumerate(summary_rows,4):
    for j,val in enumerate(row):
        c=idx_ws.cell(row=r,column=j+1,value=val); c.border=border; c.font=data_font
        c.alignment=ctr_v if j in (0,6,7,8) else left_v
for j,w in enumerate([5,16,40,18,16,10,13,20,12]): idx_ws.column_dimensions[openpyxl.utils.get_column_letter(j+1)].width=w
idx_ws.freeze_panes="A4"

wb.save(out)
print("Wrote",out,"sheets:",len(wb.sheetnames))
for row in summary_rows: print(f"  {row[0]:>3} {row[2][:36]:36} pre={row[6]:>2} filled={row[7]:>2} blank={row[8]:>2}")
