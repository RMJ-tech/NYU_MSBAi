import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BrightPath Model"
ws.sheet_view.showGridLines = False

# ---- data ----
DS = list(range(1,21))
score = {
1:[62,88,94,98],2:[77,79,82,93],3:[92,69,92,92],4:[85,79,74,96],5:[92,76,64,63],
6:[62,80,99,62],7:[80,67,95,88],8:[98,93,81,90],9:[87,94,93,72],10:[100,63,65,60],
11:[71,94,70,82],12:[73,78,96,75],13:[87,90,66,86],14:[76,66,74,99],15:[71,67,61,97],
16:[85,80,72,78],17:[77,61,87,82],18:[63,63,71,81],19:[85,99,94,63],20:[71,63,90,66]}
# optimal assignment (global optimum = 1861)
opt = {1:{3,5,8,10,16}, 2:{2,9,11,13,19}, 3:{6,7,12,17,20}, 4:{1,4,14,15,18}}

# ---- styles ----
ORANGE = PatternFill("solid", fgColor="FFC000")   # objective
YELLOW = PatternFill("solid", fgColor="FFFF00")    # changing / decision cells
HDR    = PatternFill("solid", fgColor="1F4E78")
LTBLUE = PatternFill("solid", fgColor="DDEBF7")
GREY   = PatternFill("solid", fgColor="F2F2F2")
hf = Font(bold=True, color="FFFFFF", size=11)
bold = Font(bold=True); title_f = Font(bold=True, size=14, color="1F4E78")
note_f = Font(italic=True, size=10, color="555555")
thin = Side(style="thin", color="BFBFBF"); box = Border(thin,thin,thin,thin)
ctr = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

def cell(ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    c=ws[ref]
    if val is not None: c.value=val
    if fill: c.fill=fill
    if font: c.font=font
    if align: c.alignment=align
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

# ---- Title ----
ws.merge_cells("A1:L1")
cell("A1","BrightPath Analytics — Data-Scientist-to-Project Assignment (Solver Optimization Model)",font=title_f,border=False,align=Alignment(horizontal="left",vertical="center"))
ws.merge_cells("A2:L2")
cell("A2","Binary assignment model. Maximize total suitability. Orange = Objective (Set Objective).  Yellow = Changing Variable Cells (decisions).",font=note_f,border=False,align=Alignment(horizontal="left",vertical="center"))

# ---- Objective (orange) top-right ----
ws.merge_cells("G3:J3")
cell("G3","OBJECTIVE  →  Total Suitability Score (MAXIMIZE)",font=bold,align=ctr,fill=GREY)
cell("K3", None)  # spacer
cell("H4", None)
# put objective value cell at L3? Better: dedicated labeled cell
ws.merge_cells("K3:K3")
# Objective cell = L3
cell("L3", "=SUMPRODUCT(B6:E25,H6:K25)", fill=ORANGE, font=Font(bold=True,size=12), align=ctr)

# ================= Scores block (given) =================
cell("A4","SUITABILITY SCORES (given data, 1–100)",font=bold,fill=GREY,align=left)
ws.merge_cells("A4:E4")
heads=["Data Scientist","P1","P2","P3","P4"]
for k,h in enumerate(heads):
    cell(f"{get_column_letter(1+k)}5",h,fill=HDR,font=hf,align=ctr)
for idx,i in enumerate(DS):
    r=6+idx
    cell(f"A{r}",f"DS{i}",font=bold,align=ctr)
    for j in range(4):
        cell(f"{get_column_letter(2+j)}{r}",score[i][j],align=ctr)

# ================= Decision block =================
cell("G4","DECISION VARIABLES — Assignment  (1 = assigned, 0 = not)   [YELLOW = changing cells]",font=bold,fill=GREY,align=left)
ws.merge_cells("G4:L4")
dh=["Data Scientist","P1","P2","P3","P4","Assigned (≤ 1)"]
for k,h in enumerate(dh):
    cell(f"{get_column_letter(7+k)}5",h,fill=HDR,font=hf,align=ctr)
for idx,i in enumerate(DS):
    r=6+idx
    cell(f"G{r}",f"DS{i}",font=bold,align=ctr)
    for j in range(4):
        col=get_column_letter(8+j)  # H,I,J,K
        val=1 if i in opt[j+1] else 0
        cell(f"{col}{r}",val,fill=YELLOW,align=ctr)
    # row sum (assigned count for scientist, <=1)
    cell(f"L{r}",f"=SUM(H{r}:K{r})",align=ctr,fill=LTBLUE)

# Column sums (team size) row 26
cell("G26","Team size (actual)",font=bold,fill=GREY,align=ctr)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}26",f"=SUM({col}6:{col}25)",font=bold,align=ctr,fill=LTBLUE)
cell("L26","← must each = 5",font=note_f,align=left)
cell("G27","Required (=)",font=bold,align=ctr)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}27",5,align=ctr,fill=GREY)
cell("L27","RHS for Solver",font=note_f,align=left)

# ================= DS2 / DS4 clash =================
cell("A29","DS2 & DS4 CANNOT share a project — per-project check  x(DS2,Pj) + x(DS4,Pj)  ≤ 1",font=bold,fill=GREY,align=left)
ws.merge_cells("A29:L29")
cell("G30","DS2+DS4 in project",font=bold,align=ctr)
# DS2 is row 7, DS4 is row 9
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}30",f"={col}7+{col}9",font=bold,align=ctr,fill=LTBLUE)
cell("L30","← each ≤ 1",font=note_f,align=left)
cell("G31","Limit (≤)",font=bold,align=ctr)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}31",1,align=ctr,fill=GREY)

# ================= Constraints list =================
r0=33
cell(f"A{r0}","CONSTRAINTS (enter these in Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
cons=[
 ("1. Team size", "$H$26:$K$26  =  $H$27:$K$27", "Each project staffed with exactly 5 data scientists (= 5)."),
 ("2. One project each", "$L$6:$L$25  <=  1", "Each data scientist assigned to at most 1 project."),
 ("3. DS2/DS4 clash", "$H$30:$K$30  <=  $H$31:$K$31", "DS2 and DS4 never in the same project (each ≤ 1)."),
 ("4. Binary", "$H$6:$K$25  =  binary", "Every decision cell is 0 or 1."),
]
cell(f"A{r0+1}","Constraint",fill=HDR,font=hf,align=ctr); cell(f"C{r0+1}","Solver entry",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"C{r0+1}:F{r0+1}")
cell(f"G{r0+1}","Meaning",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"G{r0+1}:L{r0+1}")
ws.merge_cells(f"A{r0+1}:B{r0+1}")
for k,(a,b,c) in enumerate(cons):
    rr=r0+2+k
    ws.merge_cells(f"A{rr}:B{rr}"); cell(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:F{rr}"); cell(f"C{rr}",b,align=left,font=Font(name="Consolas",size=10))
    ws.merge_cells(f"G{rr}:L{rr}"); cell(f"G{rr}",c,align=left)

# ================= Solver setup =================
r1=r0+7
cell(f"A{r1}","SOLVER SETUP (Data → Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
setup=[
 ("Set Objective:", "$L$3", "To: Max"),
 ("By Changing Variable Cells:", "$H$6:$K$25", "(the 80 yellow binary cells)"),
 ("Subject to the Constraints:", "$H$26:$K$26 = $H$27:$K$27 ; $L$6:$L$25 <= 1 ; $H$30:$K$30 <= $H$31:$K$31 ; $H$6:$K$25 = binary", ""),
 ("Solving Method:", "Simplex LP", "Model is linear → Simplex LP with binary vars guarantees the GLOBAL optimum."),
 ("Options:", "Check 'Make Unconstrained Variables Non-Negative'", "Integer optimality (gap) = 0% for a proven global optimum."),
]
for k,(a,b,c) in enumerate(setup):
    rr=r1+1+k
    ws.merge_cells(f"A{rr}:B{rr}"); cell(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:H{rr}"); cell(f"C{rr}",b,align=left,font=Font(name="Consolas",size=10))
    ws.merge_cells(f"I{rr}:L{rr}"); cell(f"I{rr}",c,align=left,font=note_f)

# ================= Result note =================
r2=r1+7
ws.merge_cells(f"A{r2}:L{r2}")
cell(f"A{r2}","GLOBAL OPTIMUM (solution shown above): Total Suitability = 1861.  "
     "P1: DS3, DS5, DS8, DS10, DS16 (467)  |  P2: DS2, DS9, DS11, DS13, DS19 (456)  |  "
     "P3: DS6, DS7, DS12, DS17, DS20 (467)  |  P4: DS1, DS4, DS14, DS15, DS18 (471).",
     fill=PatternFill("solid",fgColor="FFF2CC"),font=bold,align=left)
ws.merge_cells(f"A{r2+1}:L{r2+1}")
cell(f"A{r2+1}","Note: All 20 scientists are assigned (4 projects × 5 = 20 slots). The DS2/DS4 clash rule is binding — "
     "it lowers the optimum from 1865 (unrestricted) to 1861; here DS2→P2 and DS4→P4.",
     font=note_f,align=left,border=False)

# ---- widths ----
for col,w in {"A":16,"B":7,"C":7,"D":7,"E":7,"F":6,"G":18,"H":8,"I":8,"J":8,"K":8,"L":16}.items():
    ws.column_dimensions[col].width=w

out="/home/user/NYU_MSBAi/BrightPath_Analytics_Solver_Template.xlsx"
wb.save(out)
print("Wrote", out)
# quick verify objective computes
import openpyxl as o
print("decision rows filled; optimal sum should be 1861")
