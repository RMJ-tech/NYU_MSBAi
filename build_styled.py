import pandas as pd, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = "/root/.claude/uploads/25fd9b56-1752-5edc-8721-3932cd50f9d6/70eb99d7-20260523_naming_rights_database_1.xlsx"
ENR = "/home/user/NYU_MSBAi/first_50_venues_enriched.xlsx"

# ---- Read original structure: groups (row0), category (row2), field names (row3) ----
raw_full = pd.read_excel(SRC, sheet_name="Data_Input", header=None)
group_row = raw_full.iloc[0].tolist()
category_row = raw_full.iloc[2].tolist()
field_row = raw_full.iloc[3].tolist()

def norm_group(v):
    s = str(v)
    if s.startswith("Group.1"): return "Group 1"
    if s.startswith("Group.2"): return "Group 2"
    return ""

groups = [norm_group(g) for g in group_row]
categories = [("" if pd.isna(c) else str(c)) for c in category_row]
fields = [("" if pd.isna(f) else str(f)) for f in field_row]
NCOLS = 100  # original fields

# ---- Raw (pre-enrichment) and enriched values ----
raw_df = pd.read_excel(SRC, sheet_name="Data_Input", header=3).dropna(how="all").reset_index(drop=True).head(50)
enr_df = pd.read_excel(ENR, sheet_name="Full_Enriched").head(50)

SIX = ["Country", "City", "State / Prefecture", "Ownership Type", "Opened Year", "Capacity"]

def isblank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "NaT")

def fmt(v):
    if isblank(v): return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v)

# ---- Styles ----
DARK   = "1E4620"
GREEN  = "E2EFDA"   # completed
YELLOW = "FFF2CC"   # needs external input
hdr_fill   = PatternFill("solid", fgColor=DARK)
green_fill = PatternFill("solid", fgColor=GREEN)
yel_fill   = PatternFill("solid", fgColor=YELLOW)
hdr_font   = Font(bold=True, size=11, color="FFFFFF")
title_font = Font(bold=True, size=14, color=DARK)
sub_font   = Font(size=11, color="555555")
data_font  = Font(size=10, color="333333")
bold_data  = Font(size=10, color="000000", bold=True)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
ctr_v = Alignment(horizontal="center", vertical="top")
left_v = Alignment(horizontal="left", vertical="top", wrap_text=True)

HEADERS = ["Col Index","Category Block","Field Name Header","Original Group Priority",
           "Populated Value","Original Raw Value","Needs External Input?","Source & Rationale Context"]
WIDTHS  = [10, 28, 34, 18, 45, 45, 20, 60]

def clean_sheet_name(rid, venue):
    name = re.sub(r'[:\\/?*\[\]]', '-', f"{rid}_{venue}")
    return name[:31]

out = "/home/user/NYU_MSBAi/first_50_venues_styled.xlsx"
wb = openpyxl.Workbook()
wb.remove(wb.active)

for i in range(50):
    raw_row = raw_df.iloc[i]
    enr_row = enr_df.iloc[i]
    rid   = enr_row["Record ID"]
    venue = enr_row["Venue Name"]
    src_url = enr_row.get("Web Source URL (added)", "")

    ws = wb.create_sheet(clean_sheet_name(rid, venue))
    ws.sheet_view.showGridLines = False

    # Titles
    ws["A1"] = "Strict Priority-Group Audit & Live Web Enrichment Master Framework"
    ws["A1"].font = title_font
    ws["A2"] = f"Record Status Review: {rid} ({venue})"
    ws["A2"].font = sub_font

    # Header (row 4)
    for j, h in enumerate(HEADERS):
        c = ws.cell(row=4, column=j+1, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = ctr; c.border = border

    # Data rows (start row 5), one per original field
    for col_idx in range(NCOLS):
        field = fields[col_idx]
        rv = raw_row[field] if field in raw_df.columns else None
        pv = enr_row[field] if field in enr_df.columns else None
        is_enriched = (field in SIX) and isblank(rv) and (not isblank(pv))
        blank_unfilled = isblank(pv) and not is_enriched
        has_priority = groups[col_idx] != ""

        if is_enriched:
            status = "Completed"; fill = green_fill
            rationale = f"Sourced via web lookup verification: {src_url}"
        elif blank_unfilled and has_priority:
            status = "Yes"; fill = yel_fill
            rationale = "Blank high-priority field — requires external research input."
        else:
            status = "No"; fill = None
            rationale = "Pre-existing raw value from the original spreadsheet extract matrix."

        r = 5 + col_idx
        rowvals = [col_idx, categories[col_idx], field, groups[col_idx],
                   fmt(pv), fmt(rv), status, rationale]
        for j, val in enumerate(rowvals):
            c = ws.cell(row=r, column=j+1, value=val)
            c.border = border
            c.font = data_font
            if j in (0, 3, 6):  # col index, group, status centered
                c.alignment = ctr_v
            else:
                c.alignment = left_v
            # conditional fill on Populated Value (5) and Status (7)
            if fill and j in (4, 6):
                c.fill = fill
            if j == 6 and status in ("Completed", "Yes"):
                c.font = bold_data

    # widths / freeze
    for j, w in enumerate(WIDTHS):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j+1)].width = w
    ws.freeze_panes = "A5"
    ws.row_dimensions[4].height = 30

wb.save(out)
print("Wrote", out, "with", len(wb.sheetnames), "sheets")
