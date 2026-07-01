import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook(); ws=wb.active; ws.title="Restaurant Model"
ws.sheet_view.showGridLines=False

sites=list(range(1,11))
RS={1:11.8,2:13.3,3:19.0,4:17.8,5:10.0,6:16.1,7:13.3,8:18.8,9:17.2,10:14.4}
OG={1:16.2,2:13.8,3:14.6,4:12.4,5:13.7,6:19.0,7:10.8,8:15.2,9:15.9,10:16.8}
edges=[(1,2),(1,3),(1,4),(1,5),(2,3),(2,4),(2,5),(3,4),(3,5),(5,8),(5,9),(6,7),(7,8),(7,9),(8,10),(9,10)]
# optimal solution (NPV 124.8): site -> 'RS'/'OG'/None
opt={3:'RS',4:'OG',5:'OG',6:'RS',7:'OG',8:'RS',9:'RS',10:'OG'}

ORANGE=PatternFill("solid",fgColor="FFC000"); YELLOW=PatternFill("solid",fgColor="FFFF00")
HDR=PatternFill("solid",fgColor="1F4E78"); LT=PatternFill("solid",fgColor="DDEBF7"); GREY=PatternFill("solid",fgColor="F2F2F2")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note_f=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=10)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)

def cell(ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    c=ws[ref]
    if val is not None: c.value=val
    if fill: c.fill=fill
    if font: c.font=font
    if align: c.alignment=align
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

ws.merge_cells("A1:L1")
cell("A1","Darten Restaurant — Site Location (Solver Optimization Model)",font=title_f,border=False,align=Alignment(horizontal="left",vertical="center"))
ws.merge_cells("A2:F2")
cell("A2","At each site build a Red Snapper, an Olive Grove, or nothing — maximize total NPV. Orange = Objective. Yellow = Changing Variable Cells (decisions).",font=note_f,border=False,align=left)

# Objective (orange), top-right
ws.merge_cells("H2:I2"); cell("H2","OBJECTIVE  →  Total NPV ($M) — MAXIMIZE",font=bold,align=ctr,fill=GREY)
cell("J2","=SUMPRODUCT(B6:B15,D6:D15)+SUMPRODUCT(C6:C15,E6:E15)",fill=ORANGE,font=Font(bold=True,size=12),align=ctr,fmt="0.0")

# ---- Data + decisions ----
cell("A4","SITE DATA & DECISIONS",font=bold,fill=GREY,align=left); ws.merge_cells("A4:F4")
heads=["Site","Red Snapper NPV ($M)","Olive Grove NPV ($M)","Build Red Snapper (0/1)","Build Olive Grove (0/1)","Restaurants at site (≤ 1)"]
for k,h in enumerate(heads):
    cell(f"{get_column_letter(1+k)}5",h,fill=HDR,font=hf,align=Alignment(horizontal="center",vertical="center",wrap_text=True))
for idx,i in enumerate(sites):
    r=6+idx
    cell(f"A{r}",i,font=bold,align=ctr)
    cell(f"B{r}",RS[i],align=ctr,fmt="0.0")
    cell(f"C{r}",OG[i],align=ctr,fmt="0.0")
    cell(f"D{r}",1 if opt.get(i)=='RS' else 0,fill=YELLOW,align=ctr)
    cell(f"E{r}",1 if opt.get(i)=='OG' else 0,fill=YELLOW,align=ctr)
    cell(f"F{r}",f"=D{r}+E{r}",fill=LT,align=ctr)

# ---- Adjacency (same-chain) table ----
cell("H4","WITHIN-15-MILE PAIRS  (no two of the SAME chain in a pair)",font=bold,fill=GREY,align=left); ws.merge_cells("H4:K4")
ah=["Site i","Site j","RS_i + RS_j (≤ 1)","OG_i + OG_j (≤ 1)"]
for k,h in enumerate(ah):
    cell(f"{get_column_letter(8+k)}5",h,fill=HDR,font=hf,align=Alignment(horizontal="center",vertical="center",wrap_text=True))
for k,(i,j) in enumerate(edges):
    r=6+k
    ri,rj=5+i,5+j
    cell(f"H{r}",i,align=ctr); cell(f"I{r}",j,align=ctr)
    cell(f"J{r}",f"=D{ri}+D{rj}",fill=LT,align=ctr)
    cell(f"K{r}",f"=E{ri}+E{rj}",fill=LT,align=ctr)

# ---- Constraints list ----
r0=25
cell(f"A{r0}","CONSTRAINTS (enter these in Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
cell(f"A{r0+1}","Constraint",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"A{r0+1}:B{r0+1}")
cell(f"C{r0+1}","Solver entry",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"C{r0+1}:F{r0+1}")
cell(f"G{r0+1}","Meaning",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"G{r0+1}:L{r0+1}")
cons=[
 ("1. One per site","$F$6:$F$15  <=  1","At most one restaurant (Red Snapper OR Olive Grove) per site."),
 ("2. Red Snapper spacing","$J$6:$J$21  <=  1","No two Red Snappers within 15 miles of each other."),
 ("3. Olive Grove spacing","$K$6:$K$21  <=  1","No two Olive Groves within 15 miles of each other."),
 ("4. Binary","$D$6:$E$15  =  binary","Each build decision is 0 or 1."),
]
for k,(a,b,c) in enumerate(cons):
    rr=r0+2+k
    ws.merge_cells(f"A{rr}:B{rr}"); cell(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:F{rr}"); cell(f"C{rr}",b,align=left,font=mono)
    ws.merge_cells(f"G{rr}:L{rr}"); cell(f"G{rr}",c,align=left)

# ---- Solver setup ----
r1=r0+7
cell(f"A{r1}","SOLVER SETUP (Data → Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
setup=[
 ("Set Objective:","$J$2","To: Max"),
 ("By Changing Variable Cells:","$D$6:$E$15","(the 20 yellow binary cells)"),
 ("Subject to the Constraints:","$F$6:$F$15 <= 1 ; $J$6:$J$21 <= 1 ; $K$6:$K$21 <= 1 ; $D$6:$E$15 = binary",""),
 ("Solving Method:","Simplex LP","Model is linear → Simplex LP with binary vars guarantees the GLOBAL optimum."),
 ("Options:","Check 'Make Unconstrained Variables Non-Negative'; Integer optimality (gap) = 0%",""),
]
for k,(a,b,c) in enumerate(setup):
    rr=r1+1+k
    ws.merge_cells(f"A{rr}:B{rr}"); cell(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:H{rr}"); cell(f"C{rr}",b,align=left,font=mono)
    ws.merge_cells(f"I{rr}:L{rr}"); cell(f"I{rr}",c,align=left,font=note_f)

# ---- Result note ----
r2=r1+7
ws.merge_cells(f"A{r2}:L{r2}")
cell(f"A{r2}","GLOBAL OPTIMUM (solution shown above): Total NPV = $124.8M.  "
     "Red Snapper at sites 3, 6, 8, 9 ; Olive Grove at sites 4, 5, 7, 10 ; sites 1 & 2 left empty.",
     fill=PatternFill("solid",fgColor="FFF2CC"),font=bold,align=left)
ws.merge_cells(f"A{r2+1}:L{r2+1}")
cell(f"A{r2+1}","Note: 'within 15 miles' is treated as symmetric (union of the table's listed pairs → 16 pairs). "
     "Different chains may sit within 15 miles of each other; only same-chain pairs are restricted.",
     font=note_f,align=left,border=False)

widths={"A":6,"B":13,"C":13,"D":13,"E":13,"F":13,"G":6,"H":7,"I":7,"J":15,"K":15,"L":10}
for col,w in widths.items(): ws.column_dimensions[col].width=w
# wrap header row height
ws.row_dimensions[5].height=42

out="/home/user/NYU_MSBAi/Restaurant_Location_Solver_Template.xlsx"
wb.save(out); print("Wrote",out)
