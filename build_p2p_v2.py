import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pulp

# ============ FULL 10-bucket model: each grade A-E x term 36/60 ============
# (name, grade, term, net_return%, exp_loss%, stress_uplift%, avg_life, supply_cap)
raw=[
 ("A-36","A",36, 6.8, 0.9, 0.5, 1.9, 0.15),
 ("A-60","A",60, 7.1, 0.9, 0.5, 3.0, 1.00),
 ("B-36","B",36, 8.3, 2.2, 0.8, 1.9, 0.20),
 ("B-60","B",60, 8.8, 2.2, 0.8, 3.0, 1.00),
 ("C-36","C",36,10.0, 4.7, 1.5, 2.0, 0.20),
 ("C-60","C",60,10.7, 4.7, 1.5, 3.2, 0.30),
 ("D-36","D",36,11.6, 7.9, 2.5, 2.1, 1.00),
 ("D-60","D",60,12.2, 7.9, 2.5, 3.4, 0.25),
 ("E-36","E",36,12.6,11.5, 3.5, 2.1, 1.00),
 ("E-60","E",60,13.3,11.5, 3.5, 3.4, 0.12),
]
# stress loss = expected loss + stress uplift
buckets=[(nm,gr,tm,ret,el,round(el+up,4),al,cap) for (nm,gr,tm,ret,el,up,al,cap) in raw]
N=len(buckets)

# ---- Solve the LP to fill the optimal weights (40% 5-yr cap ENFORCED) ----
ret=[b[3] for b in buckets]; el=[b[4] for b in buckets]; sl=[b[5] for b in buckets]
al=[b[6] for b in buckets]; cap=[b[7] for b in buckets]; term=[b[2] for b in buckets]; grade=[b[1] for b in buckets]
fy=[1 if t==60 else 0 for t in term]
m=pulp.LpProblem("p2p",pulp.LpMaximize)
w=[pulp.LpVariable(f"w{i}",0,cap[i]) for i in range(N)]
m+=pulp.lpSum(ret[i]*w[i] for i in range(N))
m+=pulp.lpSum(w)==1
m+=pulp.lpSum(el[i]*w[i] for i in range(N))<=5
m+=pulp.lpSum(sl[i]*w[i] for i in range(N))<=7
m+=pulp.lpSum(al[i]*w[i] for i in range(N))<=2.6
m+=pulp.lpSum(fy[i]*w[i] for i in range(N))>=0.20
m+=pulp.lpSum(fy[i]*w[i] for i in range(N))<=0.40
m+=pulp.lpSum(w[i] for i in range(N) if grade[i]=="A")>=0.15
m+=pulp.lpSum(w[i] for i in range(N) if grade[i]=="C")<=0.35
m+=pulp.lpSum(w[i] for i in range(N) if grade[i]=="D")<=0.25
m+=pulp.lpSum(w[i] for i in range(N) if grade[i]=="E")<=0.15
m.solve(pulp.PULP_CBC_CMD(msg=0))
wopt=[w[i].value() for i in range(N)]
print("STATUS:",pulp.LpStatus[m.status],"RETURN=%.4f%%"%pulp.value(m.objective))

# ============ Build the workbook ============
wb=openpyxl.Workbook(); ws=wb.active; ws.title="OpenCircle Model"
ws.sheet_view.showGridLines=False
ORANGE=PatternFill("solid",fgColor="FFC000"); YELLOW=PatternFill("solid",fgColor="FFFF00")
HDR=PatternFill("solid",fgColor="1F4E78"); INP=PatternFill("solid",fgColor="DDEBF7")
LHS_F=PatternFill("solid",fgColor="E2EFDA"); RHS_F=PatternFill("solid",fgColor="FCE4D6")
SIGN_F=PatternFill("solid",fgColor="FFF2CC"); GREY=PatternFill("solid",fgColor="F2F2F2")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note_f=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=10)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)
ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)

def c(ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    cell=ws[ref]
    if val is not None: cell.value=val
    if fill: cell.fill=fill
    if font: cell.font=font
    if align: cell.alignment=align
    if border: cell.border=box
    if fmt: cell.number_format=fmt
    return cell

ws.merge_cells("A1:I1")
c("A1","OpenCircle Auto-Invest — Portfolio Allocation (Solver Model)  [FULL 10-bucket, 5-yr cap 40% enforced]",font=title_f,border=False,align=Alignment(horizontal="left",vertical="center"))
ws.merge_cells("A2:I2")
c("A2","Orange = Objective.  Yellow = changing cells (weights).  Blue = inputs.  Green = constraint LHS.  Peach = RHS.  Stress loss = expected loss + stress uplift.",font=note_f,border=False,align=left)

# ---- Input + decision table (rows 5..14) ----
heads=["Bucket","Grade","Term (m)","Net Return (%)","Exp. Loss (%/yr)","Stress Loss (%/yr)","Avg Life (yrs)","Supply Cap","Weight (decision)"]
for k,h in enumerate(heads):
    c(f"{get_column_letter(1+k)}4",h,fill=HDR,font=hf,align=ctrw)
first,last=5,5+N-1
for idx,(nm,gr,tm,r_,el_,sl_,al_,cap_) in enumerate(buckets):
    r=5+idx
    c(f"A{r}",nm,font=bold,align=ctr); c(f"B{r}",gr,align=ctr); c(f"C{r}",tm,align=ctr)
    c(f"D{r}",r_,fill=INP,align=ctr,fmt="0.0")
    c(f"E{r}",el_,fill=INP,align=ctr,fmt="0.0")
    c(f"F{r}",sl_,fill=INP,align=ctr,fmt="0.0")
    c(f"G{r}",al_,fill=INP,align=ctr,fmt="0.0")
    c(f"H{r}",cap_,fill=INP,align=ctr,fmt="0%")
    c(f"I{r}",round(wopt[idx],6),fill=YELLOW,align=ctr,fmt="0.00%")   # optimal weight

# ---- Objective ----
OBJ=16
c(f"A{OBJ}","OBJECTIVE (Max) -> Portfolio Net Return (%)",font=bold,fill=GREY,align=left); ws.merge_cells(f"A{OBJ}:B{OBJ}")
c(f"C{OBJ}",f"=SUMPRODUCT(D{first}:D{last},I{first}:I{last})",fill=ORANGE,font=Font(bold=True,size=12),align=ctr,fmt="0.000")
c(f"D{OBJ}","optimum = 10.094%",font=note_f,border=False,align=left); ws.merge_cells(f"D{OBJ}:I{OBJ}")

# ---- Constraints ----
CH=18
c(f"A{CH}","CONSTRAINTS  (LHS formula  |  sign  |  RHS)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
c(f"A{CH+1}","Constraint",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"A{CH+1}:B{CH+1}")
c(f"C{CH+1}","LHS (formula cell)",fill=HDR,font=hf,align=ctr)
c(f"D{CH+1}","sign",fill=HDR,font=hf,align=ctr)
c(f"E{CH+1}","RHS",fill=HDR,font=hf,align=ctr)
c(f"F{CH+1}","Note",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"F{CH+1}:I{CH+1}")
def rng_grade(g):
    rows=[5+i for i in range(N) if grade[i]==g]
    return "+".join(f"I{r}" for r in rows)
cons=[
 ("1. Fully invested",f"=SUM(I{first}:I{last})","=",1,"0%","weights sum to 100%"),
 ("2. Loss - normal",f"=SUMPRODUCT(E{first}:E{last},I{first}:I{last})","<=",5,"0.00","portfolio expected loss <= 5%/yr"),
 ("3. Loss - stress",f"=SUMPRODUCT(F{first}:F{last},I{first}:I{last})","<=",7,"0.00","stressed loss <= 7%/yr"),
 ("4. Avg life",f"=SUMPRODUCT(G{first}:G{last},I{first}:I{last})","<=",2.6,"0.000","weighted avg life <= 2.6 yrs"),
 ("5. 5-yr min","=SUMIF($C$5:$C$14,60,I5:I14)",">=",0.20,"0%","five-year (60m) share >= 20%"),
 ("6. 5-yr max","=SUMIF($C$5:$C$14,60,I5:I14)","<=",0.40,"0%","five-year (60m) share <= 40%"),
 ("7. Grade A min",f"={rng_grade('A')}",">=",0.15,"0%","Grade A (A-36+A-60) >= 15%"),
 ("8. Grade C max",f"={rng_grade('C')}","<=",0.35,"0%","Grade C (C-36+C-60) <= 35%"),
 ("9. Grade D max",f"={rng_grade('D')}","<=",0.25,"0%","Grade D (D-36+D-60) <= 25%"),
 ("10. Grade E max",f"={rng_grade('E')}","<=",0.15,"0%","Grade E (E-36+E-60) <= 15%"),
]
r=CH+2
for lab,lhs,sign,rhs,fmt,note in cons:
    ws.merge_cells(f"A{r}:B{r}"); c(f"A{r}",lab,font=bold,align=left)
    c(f"C{r}",lhs,fill=LHS_F,align=ctr,font=mono,fmt=fmt)
    c(f"D{r}",sign,fill=SIGN_F,align=ctr)
    c(f"E{r}",rhs,fill=RHS_F,align=ctr,fmt=fmt)
    ws.merge_cells(f"F{r}:I{r}"); c(f"F{r}",note,align=left,font=note_f)
    r+=1
ws.merge_cells(f"A{r}:B{r}"); c(f"A{r}","11. Supply caps",font=bold,align=left)
c(f"C{r}",f"I{first}:I{last}",fill=LHS_F,align=ctr,font=mono)
c(f"D{r}","<=",fill=SIGN_F,align=ctr)
c(f"E{r}",f"H{first}:H{last}",fill=RHS_F,align=ctr,font=mono)
ws.merge_cells(f"F{r}:I{r}"); c(f"F{r}","each bucket <= its supply cap (column H)",align=left,font=note_f)
supply_row=r

# ---- Solver setup ----
r0=r+2
c(f"A{r0}","SOLVER SETUP (Data -> Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
cr=CH+2  # first constraint row = 20
setup=[
 ("Set Objective",f"$C${OBJ}","To: Max"),
 ("By Changing Variable Cells",f"$I${first}:$I${last}","the 10 yellow weights"),
 ("Subject to the Constraints",
  f"$C${cr}=$E${cr} ; $C${cr+1}<=$E${cr+1} ; $C${cr+2}<=$E${cr+2} ; $C${cr+3}<=$E${cr+3} ; "
  f"$C${cr+4}>=$E${cr+4} ; $C${cr+5}<=$E${cr+5} ; $C${cr+6}>=$E${cr+6} ; $C${cr+7}<=$E${cr+7} ; "
  f"$C${cr+8}<=$E${cr+8} ; $C${cr+9}<=$E${cr+9} ; $I${first}:$I${last}<=$H${first}:$H${last}",""),
 ("Options","Check 'Make Unconstrained Variables Non-Negative'",""),
 ("Solving Method","Simplex LP","linear -> GLOBAL optimum"),
]
c(f"A{r0+1}","Item",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"A{r0+1}:B{r0+1}")
c(f"C{r0+1}","Entry",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"C{r0+1}:G{r0+1}")
c(f"H{r0+1}","Note",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"H{r0+1}:I{r0+1}")
for k,(a,b,d) in enumerate(setup):
    rr=r0+2+k
    ws.merge_cells(f"A{rr}:B{rr}"); c(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:G{rr}"); c(f"C{rr}",b,align=left,font=mono)
    ws.merge_cells(f"H{rr}:I{rr}"); c(f"H{rr}",d,align=left,font=note_f)

# ---- Result banner ----
rb=r0+2+len(setup)+1
ws.merge_cells(f"A{rb}:I{rb}")
c(f"A{rb}","GLOBAL OPTIMUM = 10.094%.  A-36 15% | B-36 17.53% | C-36 5% | C-60 30% | D-36 22.47% | D-60 2.53% | E-60 7.47%.  "
  "5-yr share = 40% (binding).  Full 10-bucket universe (grades A-E x 36/60m).",
  fill=PatternFill("solid",fgColor="FFF2CC"),font=bold,align=left)

for col,w_ in {"A":16,"B":7,"C":24,"D":6,"E":9,"F":10,"G":11,"H":11,"I":16}.items():
    ws.column_dimensions[col].width=w_
ws.row_dimensions[4].height=42

out="/home/user/NYU_MSBAi/P2P_OpenCircle_Solver_Template_v2.xlsx"
wb.save(out); print("Wrote",out)
