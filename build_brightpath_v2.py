import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook(); ws=wb.active; ws.title="BrightPath Model"
ws.sheet_view.showGridLines=False

DS=list(range(1,21))
score={1:[62,88,94,98],2:[77,79,82,93],3:[92,69,92,92],4:[85,79,74,96],5:[92,76,64,63],
6:[62,80,99,62],7:[80,67,95,88],8:[98,93,81,90],9:[87,94,93,72],10:[100,63,65,60],
11:[71,94,70,82],12:[73,78,96,75],13:[87,90,66,86],14:[76,66,74,99],15:[71,67,61,97],
16:[85,80,72,78],17:[77,61,87,82],18:[63,63,71,81],19:[85,99,94,63],20:[71,63,90,66]}
opt={1:{3,5,8,10,16},2:{2,9,11,13,19},3:{6,7,12,17,20},4:{1,4,14,15,18}}

ORANGE=PatternFill("solid",fgColor="FFC000"); YELLOW=PatternFill("solid",fgColor="FFFF00")
HDR=PatternFill("solid",fgColor="1F4E78"); LHS_F=PatternFill("solid",fgColor="DDEBF7")
RHS_F=PatternFill("solid",fgColor="FCE4D6"); GREY=PatternFill("solid",fgColor="F2F2F2")
SIGN_F=PatternFill("solid",fgColor="FFF2CC")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note_f=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=10)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)
ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)

def cell(ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    c=ws[ref]
    if val is not None: c.value=val
    if fill: c.fill=fill
    if font: c.font=font
    if align: c.alignment=align
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

ws.merge_cells("A1:N1")
cell("A1","BrightPath Analytics — Data-Scientist Assignment (Solver Model, constraint-separated format)",font=title_f,border=False,align=Alignment(horizontal="left",vertical="center"))
ws.merge_cells("A2:F2")
cell("A2","Orange = Objective.  Yellow = changing cells (decisions).  Blue = constraint LHS (formula cells).  Peach = RHS limits.",font=note_f,border=False,align=left)

# Objective (orange)
ws.merge_cells("H2:J2"); cell("H2","OBJECTIVE → Total Suitability (MAX)",font=bold,align=ctr,fill=GREY)
cell("K2","=SUMPRODUCT(B6:E25,H6:K25)",fill=ORANGE,font=Font(bold=True,size=12),align=ctr)

# ---- Scores ----
cell("A4","SUITABILITY SCORES (given, 1-100)",font=bold,fill=GREY,align=left); ws.merge_cells("A4:E4")
for k,h in enumerate(["Data Scientist","P1","P2","P3","P4"]):
    cell(f"{get_column_letter(1+k)}5",h,fill=HDR,font=hf,align=ctr)
for idx,i in enumerate(DS):
    r=6+idx; cell(f"A{r}",f"DS{i}",font=bold,align=ctr)
    for j in range(4): cell(f"{get_column_letter(2+j)}{r}",score[i][j],align=ctr)

# ---- Decisions + row-sum constraint (Assigned <= 1) ----
cell("G4","DECISION VARIABLES (1=assigned)   [YELLOW]",font=bold,fill=GREY,align=left); ws.merge_cells("G4:K4")
cell("L4","1-project rule",font=bold,fill=GREY,align=ctr); ws.merge_cells("L4:N4")
for k,h in enumerate(["Data Scientist","P1","P2","P3","P4"]):
    cell(f"{get_column_letter(7+k)}5",h,fill=HDR,font=hf,align=ctr)
cell("L5","Assigned  (LHS =SUM)",fill=HDR,font=hf,align=ctrw)
cell("M5","sign",fill=HDR,font=hf,align=ctr)
cell("N5","Limit (RHS)",fill=HDR,font=hf,align=ctrw)
for idx,i in enumerate(DS):
    r=6+idx; cell(f"G{r}",f"DS{i}",font=bold,align=ctr)
    for j in range(4):
        col=get_column_letter(8+j)
        cell(f"{col}{r}",1 if i in opt[j+1] else 0,fill=YELLOW,align=ctr)
    cell(f"L{r}",f"=SUM(H{r}:K{r})",fill=LHS_F,align=ctr)      # LHS
    cell(f"M{r}","<=",fill=SIGN_F,align=ctr)                   # sign
    cell(f"N{r}",1,fill=RHS_F,align=ctr)                       # RHS

# ---- Team size constraint (= 5), per project, separated ----
ws.merge_cells("A27:F27"); cell("A27","CONSTRAINT 1 — TEAM SIZE (each project must = 5)",font=bold,fill=GREY,align=left)
for j in range(4):
    col=get_column_letter(8+j); cell(f"{col}27",f"P{j+1}",fill=HDR,font=hf,align=ctr)
cell("L27","",fill=GREY)
cell("G28","# of DS in project  (LHS =SUM)",font=bold,align=left,fill=GREY)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}28",f"=SUM({col}6:{col}25)",fill=LHS_F,font=bold,align=ctr)
cell("L28","= (each)",fill=SIGN_F,align=ctr)
cell("G29","Required team size (RHS)",font=bold,align=left,fill=GREY)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}29",5,fill=RHS_F,font=bold,align=ctr)
cell("L29","RHS = 5",fill=RHS_F,align=ctr)

# ---- DS2/DS4 conflict constraint (<=1), per project ----
ws.merge_cells("A31:F31"); cell("A31","CONSTRAINT 3 — DS2/DS4 CONFLICT (each project: DS2 + DS4 <= 1)",font=bold,fill=GREY,align=left)
for j in range(4):
    col=get_column_letter(8+j); cell(f"{col}31",f"P{j+1}",fill=HDR,font=hf,align=ctr)
cell("L31","",fill=GREY)
cell("G32","DS2 + DS4 in project  (LHS)",font=bold,align=left,fill=GREY)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}32",f"={col}7+{col}9",fill=LHS_F,font=bold,align=ctr)
cell("L32","<= (each)",fill=SIGN_F,align=ctr)
cell("G33","Limit (RHS)",font=bold,align=left,fill=GREY)
for j in range(4):
    col=get_column_letter(8+j)
    cell(f"{col}33",1,fill=RHS_F,font=bold,align=ctr)
cell("L33","RHS = 1",fill=RHS_F,align=ctr)

# ================= Solver setup (professor-style text block) =================
r0=36
cell(f"A{r0}","SOLVER SETUP  (Data → Solver)  — enter the cell ranges below",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
rows=[
 ("Decisions (Changing Cells)","$H$6:$K$25","80 binary cells (yellow)"),
 ("Objective","Max  $K$2","= SUMPRODUCT(scores, decisions)"),
 ("Constraint 1 — Team size","$H$28:$K$28  =  $H$29:$K$29","each project has exactly 5 DS"),
 ("Constraint 2 — 1 project each","$L$6:$L$25  <=  $N$6:$N$25","each DS in at most 1 project"),
 ("Constraint 3 — DS2/DS4 conflict","$H$32:$K$32  <=  $H$33:$K$33","DS2 & DS4 never in same project"),
 ("Constraint 4 — Binary","$H$6:$K$25  =  binary","each cell is 0 or 1"),
 ("Solving Method","Simplex LP","linear + binary → GLOBAL optimum"),
]
cell(f"A{r0+1}","Item",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"A{r0+1}:C{r0+1}")
cell(f"D{r0+1}","Solver entry",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"D{r0+1}:H{r0+1}")
cell(f"I{r0+1}","Meaning",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"I{r0+1}:N{r0+1}")
for k,(a,b,c) in enumerate(rows):
    rr=r0+2+k
    ws.merge_cells(f"A{rr}:C{rr}"); cell(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"D{rr}:H{rr}"); cell(f"D{rr}",b,align=left,font=mono)
    ws.merge_cells(f"I{rr}:N{rr}"); cell(f"I{rr}",c,align=left)

r1=r0+10
ws.merge_cells(f"A{r1}:N{r1}")
cell(f"A{r1}","GLOBAL OPTIMUM = 1861.  P1: DS3,5,8,10,16 (467) | P2: DS2,9,11,13,19 (456) | P3: DS6,7,12,17,20 (467) | P4: DS1,4,14,15,18 (471).",
     fill=PatternFill("solid",fgColor="FFF2CC"),font=bold,align=left)

for col,w in {"A":15,"B":6,"C":6,"D":6,"E":6,"F":4,"G":26,"H":8,"I":8,"J":8,"K":8,"L":16,"M":10,"N":12}.items():
    ws.column_dimensions[col].width=w
ws.row_dimensions[5].height=30

out="/home/user/NYU_MSBAi/BrightPath_Analytics_Solver_Template_v2.xlsx"
wb.save(out); print("Wrote",out)
