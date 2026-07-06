import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook(); ws=wb.active; ws.title="OpenCircle Model"
ws.sheet_view.showGridLines=False

# bucket: (name, grade, term, net_return%, exp_loss%, stress_loss%(=EL+uplift), avg_life, supply_cap_frac)
buckets=[
 ("A-36","A",36, 6.8, 0.9, 1.4, 1.9, 0.15),
 ("B-36","B",36, 8.3, 2.2, 3.0, 1.9, 0.20),
 ("C-36","C",36,10.0, 4.7, 6.2, 2.0, 0.20),
 ("C-60","C",60,10.7, 4.7, 6.2, 3.2, 0.30),
 ("D-60","D",60,12.2, 7.9,10.4, 3.4, 0.25),
 ("E-60","E",60,13.3,11.5,15.0, 3.4, 0.12),
]
ORANGE=PatternFill("solid",fgColor="FFC000"); YELLOW=PatternFill("solid",fgColor="FFFF00")
HDR=PatternFill("solid",fgColor="1F4E78"); INP=PatternFill("solid",fgColor="DDEBF7")
LHS_F=PatternFill("solid",fgColor="E2EFDA"); RHS_F=PatternFill("solid",fgColor="FCE4D6")
SIGN_F=PatternFill("solid",fgColor="FFF2CC"); GREY=PatternFill("solid",fgColor="F2F2F2")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note_f=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=10)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)
ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)

def c(ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    cell=ws[ref]
    if val is not None: cell.value=val
    if fill: cell.fill=fill
    if font: cell.font=font
    if align: cell.alignment=align
    if border: cell.border=box
    if fmt: cell.number_format=fmt
    return cell

ws.merge_cells("A1:I1")
c("A1","OpenCircle Auto-Invest — Portfolio Allocation (Solver Model)  [decisions blank — solve it yourself]",font=title_f,border=False,align=Alignment(horizontal="left",vertical="center"))
ws.merge_cells("A2:I2")
c("A2","Orange = Objective.  Yellow = changing cells (weights, start 0).  Blue = inputs.  Green = constraint LHS.  Peach = RHS.",font=note_f,border=False,align=left)

# ---- Input + decision table ----
heads=["Bucket","Grade","Term (m)","Net Return (%)","Exp. Loss (%/yr)","Stress Loss (%/yr)","Avg Life (yrs)","Supply Cap","Weight (decision)"]
for k,h in enumerate(heads):
    c(f"{get_column_letter(1+k)}4",h,fill=HDR,font=hf,align=ctrw)
for idx,(nm,gr,tm,ret,el,sl,al,cap) in enumerate(buckets):
    r=5+idx
    c(f"A{r}",nm,font=bold,align=ctr)
    c(f"B{r}",gr,align=ctr)
    c(f"C{r}",tm,align=ctr)
    c(f"D{r}",ret,fill=INP,align=ctr,fmt="0.0")
    c(f"E{r}",el,fill=INP,align=ctr,fmt="0.0")
    c(f"F{r}",sl,fill=INP,align=ctr,fmt="0.0")
    c(f"G{r}",al,fill=INP,align=ctr,fmt="0.0")
    c(f"H{r}",cap,fill=INP,align=ctr,fmt="0%")
    c(f"I{r}",0,fill=YELLOW,align=ctr,fmt="0.0%")     # decision, start 0

# ---- Objective (orange) ----
c("A12","OBJECTIVE (Max) → Portfolio Net Return (%)",font=bold,fill=GREY,align=left); ws.merge_cells("A12:B12")
c("C12","=SUMPRODUCT(D5:D10,I5:I10)",fill=ORANGE,font=Font(bold=True,size=12),align=ctr,fmt="0.000")
c("D12","(shows 0 until you solve)",font=note_f,border=False,align=left); ws.merge_cells("D12:I12")

# ---- Constraints (LHS | sign | RHS) ----
c("A14","CONSTRAINTS  (LHS formula  |  sign  |  RHS)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
c("A15","Constraint",fill=HDR,font=hf,align=ctr); ws.merge_cells("A15:B15")
c("C15","LHS (formula cell)",fill=HDR,font=hf,align=ctr)
c("D15","sign",fill=HDR,font=hf,align=ctr)
c("E15","RHS",fill=HDR,font=hf,align=ctr)
c("F15","Note",fill=HDR,font=hf,align=ctr); ws.merge_cells("F15:I15")
cons=[
 ("1. Fully invested","=SUM(I5:I10)","=",1,"0%","weights sum to 100%"),
 ("2. Loss — normal","=SUMPRODUCT(E5:E10,I5:I10)","<=",5,"0.0","portfolio expected loss ≤ 5%/yr"),
 ("3. Loss — stress","=SUMPRODUCT(F5:F10,I5:I10)","<=",7,"0.0","stressed loss ≤ 7%/yr"),
 ("4. Avg life","=SUMPRODUCT(G5:G10,I5:I10)","<=",2.6,"0.00","weighted avg life ≤ 2.6 yrs"),
 ("5. 5-yr min","=SUM(I8:I10)",">=",0.20,"0%","five-year (60m) share ≥ 20%"),
 ("6. 5-yr max","=SUM(I8:I10)","<=",0.40,"0%","five-year (60m) share ≤ 40%"),
 ("7. Grade A min","=I5",">=",0.15,"0%","Grade A ≥ 15% (stability)"),
 ("8. Grade C max","=I7+I8","<=",0.35,"0%","Grade C (C-36+C-60) ≤ 35%"),
 ("9. Grade D max","=I9","<=",0.25,"0%","Grade D ≤ 25%"),
 ("10. Grade E max","=I10","<=",0.15,"0%","Grade E ≤ 15%"),
]
for k,(lab,lhs,sign,rhs,fmt,note) in enumerate(cons):
    r=16+k
    ws.merge_cells(f"A{r}:B{r}"); c(f"A{r}",lab,font=bold,align=left)
    c(f"C{r}",lhs,fill=LHS_F,align=ctr,font=mono,fmt=("0.000" if fmt=="0.00" else fmt))
    c(f"D{r}",sign,fill=SIGN_F,align=ctr)
    c(f"E{r}",rhs,fill=RHS_F,align=ctr,fmt=fmt)
    ws.merge_cells(f"F{r}:I{r}"); c(f"F{r}",note,align=left,font=note_f)
# supply cap (element-wise) row
r=16+len(cons)
ws.merge_cells(f"A{r}:B{r}"); c(f"A{r}","11. Supply caps",font=bold,align=left)
c(f"C{r}","I5:I10",fill=LHS_F,align=ctr,font=mono)
c(f"D{r}","<=",fill=SIGN_F,align=ctr)
c(f"E{r}","H5:H10",fill=RHS_F,align=ctr,font=mono)
ws.merge_cells(f"F{r}:I{r}"); c(f"F{r}","each bucket ≤ its supply cap (column H)",align=left,font=note_f)

# ---- Solver setup ----
r0=r+2
c(f"A{r0}","SOLVER SETUP (Data → Solver)",font=Font(bold=True,size=12,color="1F4E78"),border=False,align=left)
setup=[
 ("Set Objective","$C$12","To: Max"),
 ("By Changing Variable Cells","$I$5:$I$10","the 6 yellow weights"),
 ("Subject to the Constraints","$C$16=$E$16 ; $C$17<=$E$17 ; $C$18<=$E$18 ; $C$19<=$E$19 ; $C$20>=$E$20 ; $C$21<=$E$21 ; $C$22>=$E$22 ; $C$23<=$E$23 ; $C$24<=$E$24 ; $C$25<=$E$25 ; $I$5:$I$10<=$H$5:$H$10",""),
 ("Options","Check 'Make Unconstrained Variables Non-Negative'",""),
 ("Solving Method","Simplex LP","linear → global optimum"),
]
c(f"A{r0+1}","Item",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"A{r0+1}:B{r0+1}")
c(f"C{r0+1}","Entry",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"C{r0+1}:G{r0+1}")
c(f"H{r0+1}","Note",fill=HDR,font=hf,align=ctr); ws.merge_cells(f"H{r0+1}:I{r0+1}")
for k,(a,b,d) in enumerate(setup):
    rr=r0+2+k
    ws.merge_cells(f"A{rr}:B{rr}"); c(f"A{rr}",a,font=bold,align=left)
    ws.merge_cells(f"C{rr}:G{rr}"); c(f"C{rr}",b,align=left,font=mono)
    ws.merge_cells(f"H{rr}:I{rr}"); c(f"H{rr}",d,align=left,font=note_f)

# NOTE: constraint rows in Solver text — careful: row indices. constraints occupy rows 16..25.
# widths
for col,w in {"A":16,"B":7,"C":24,"D":6,"E":9,"F":10,"G":11,"H":11,"I":16}.items():
    ws.column_dimensions[col].width=w
ws.row_dimensions[4].height=42

out="/home/user/NYU_MSBAi/P2P_OpenCircle_Solver_Template.xlsx"
wb.save(out); print("Wrote",out)
# sanity: constraint rows
print("constraint rows 16..25 map to cons list + supply at 26? actually supply row =",16+len(cons))
