import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook()
INP=PatternFill("solid",fgColor="DDEBF7"); ASM=PatternFill("solid",fgColor="FFFF00")
FOR=PatternFill("solid",fgColor="FFC000"); HDR=PatternFill("solid",fgColor="1F4E78")
GREY=PatternFill("solid",fgColor="F2F2F2")
hf=Font(bold=True,color="FFFFFF",size=11); bold=Font(bold=True); title_f=Font(bold=True,size=14,color="1F4E78")
note=Font(italic=True,size=10,color="555555"); mono=Font(name="Consolas",size=9)
thin=Side(style="thin",color="BFBFBF"); box=Border(thin,thin,thin,thin)
ctr=Alignment(horizontal="center",vertical="center"); left=Alignment(horizontal="left",vertical="center",wrap_text=True)
ctrw=Alignment(horizontal="center",vertical="center",wrap_text=True)

def C(ws,ref,val=None,fill=None,font=None,align=None,border=True,fmt=None):
    c=ws[ref]
    if val is not None: c.value=val
    if fill: c.fill=fill
    if font: c.font=font
    if align: c.alignment=align
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

# ==================== Sheet 4: Risk Management ====================
ws=wb.active; ws.title="4_RiskManagement"; ws.sheet_view.showGridLines=False
ws.merge_cells("A1:H1")
C(ws,"A1","4. Risk Management — Total Aggregated Loss (Crystal Ball, 10,000 trials)",font=title_f,border=False,align=left)
ws.merge_cells("A2:H2")
C(ws,"A2","Blue = inputs.  Yellow = CB Assumptions (Yes-No occurrence, Triangular loss).  Orange = Forecast (total loss).  (=CB.* shows #NAME? until CB loaded.)",font=note,border=False,align=left)

heads=["Event","Prob / yr","Most Likely ($M)","Min ($M)","Max ($M)","Occurs? (CB.YesNo)","Loss if occurs (CB.Triangular)","Event Loss ($M)"]
for k,h in enumerate(heads):
    C(ws,f"{get_column_letter(1+k)}4",h,fill=HDR,font=hf,align=ctrw)

events=[
 ("IT system, major failure",0.010,5.0,3,7),
 ("Problem with manufacturing process",0.025,1.0,0.5,1.5),
 ("Serious illness of a Board member",0.050,0.1,0.05,0.15),
 ("Employee wins law suit",0.080,2.5,2,3),
 ("Entry of new competitor",0.100,10.0,5,15),
 ("Failure of new product launch",0.075,6.0,4,8),
 ("Strengthening of $ Xrate",0.350,1.0,0.5,1.5),
 ("Fire in head office",0.020,2.5,2,3),
 ("Fraud",0.005,5.0,4,6),
 ("Confidential data lost",0.010,3.0,2,4),
 ("Large customer goes bankrupt owing money",0.020,5.0,5,5),  # always 5 if occurs
]
r0=5
for i,(nm,p,ml,mn,mx) in enumerate(events):
    r=r0+i
    C(ws,f"A{r}",nm,align=left)
    C(ws,f"B{r}",p,fill=INP,align=ctr,fmt="0.0%")
    C(ws,f"C{r}",ml,fill=INP,align=ctr,fmt="0.00")
    C(ws,f"D{r}",mn,fill=INP,align=ctr,fmt="0.00")
    C(ws,f"E{r}",mx,fill=INP,align=ctr,fmt="0.00")
    C(ws,f"F{r}",f"=CB.YesNo(B{r})",fill=ASM,align=ctr,font=mono)            # occurrence
    if mn==mx:  # last event: constant loss (no triangular)
        C(ws,f"G{r}",f"=C{r}",fill=INP,align=ctr,font=mono)
    else:
        C(ws,f"G{r}",f"=CB.Triangular(D{r},C{r},E{r})",fill=ASM,align=ctr,font=mono)  # (min,likely,max)
    C(ws,f"H{r}",f"=F{r}*G{r}",align=ctr,fmt="0.00")
rtot=r0+len(events)+1
C(ws,f"A{rtot}","TOTAL AGGREGATED LOSS  →  FORECAST",font=bold,fill=GREY,align=left); ws.merge_cells(f"A{rtot}:G{rtot}")
C(ws,f"H{rtot}",f"=SUM(H{r0}:H{r0+len(events)-1})",fill=FOR,font=Font(bold=True,size=12),align=ctr,fmt="0.00")
# how-to
r=rtot+2
C(ws,f"A{r}","HOW TO USE",font=Font(bold=True,color="1F4E78"),border=False,align=left)
for i,t in enumerate([
 "1. Define F5:F15 (occurrence) and G5:G14 (Triangular loss) as Assumptions.  (Event 11 loss is a constant $5M, not random.)",
 f"2. Define H{rtot} (Total Aggregated Loss) as the Forecast cell.",
 "3. Run Preferences -> 10,000 trials -> Run.",
 "4. Report from the Statistics/Percentiles table:  Mean = expected total loss;  Std Dev;  95th percentile (95% VaR).  Capture the frequency chart.",
]):
    C(ws,f"A{r+1+i}",t,border=False,align=left,font=note); ws.merge_cells(f"A{r+1+i}:H{r+1+i}")
for col,w in {"A":40,"B":9,"C":15,"D":9,"E":9,"F":18,"G":24,"H":14}.items(): ws.column_dimensions[col].width=w
ws.row_dimensions[4].height=42

# ==================== Sheet 5: Hometown Insurance ====================
ws2=wb.create_sheet("5_HometownInsurance"); ws2.sheet_view.showGridLines=False
ws2.merge_cells("A1:F1")
C(ws2,"A1","5. Hometown Insurance — 10-yr Annuity Profit (Crystal Ball, 10,000 trials)",font=title_f,border=False,align=left)
ws2.merge_cells("A2:F2")
C(ws2,"A2","Blue = inputs.  Yellow = CB Assumption (annual return).  Orange = Forecast (Hometown profit).",font=note,border=False,align=left)

C(ws2,"A4","Initial investment ($)",font=bold,align=left); C(ws2,"B4",500000,fill=INP,align=ctr,fmt="#,##0")
C(ws2,"A5","Mean annual return",font=bold,align=left);     C(ws2,"B5",0.08,fill=INP,align=ctr,fmt="0.0%")
C(ws2,"A6","Std dev of return",font=bold,align=left);      C(ws2,"B6",0.03,fill=INP,align=ctr,fmt="0.0%")
C(ws2,"A7","Guaranteed floor (min)",font=bold,align=left); C(ws2,"B7",0.05,fill=INP,align=ctr,fmt="0.0%")
C(ws2,"A8","Rate cap (max)",font=bold,align=left);         C(ws2,"B8",0.075,fill=INP,align=ctr,fmt="0.0%")

h2=["Year","Actual Return (CB.Normal)","Credited Rate =MIN(MAX(r,floor),cap)","Fund Value ($)","Investor Value ($)"]
for k,h in enumerate(h2):
    C(ws2,f"{get_column_letter(1+k)}10",h,fill=HDR,font=hf,align=ctrw)
# Year 0 start
C(ws2,"A11",0,align=ctr); C(ws2,"B11","—",align=ctr); C(ws2,"C11","—",align=ctr)
C(ws2,"D11","=B4",align=ctr,fmt="#,##0"); C(ws2,"E11","=B4",align=ctr,fmt="#,##0")
for y in range(1,11):
    r=11+y
    C(ws2,f"A{r}",y,align=ctr)
    C(ws2,f"B{r}","=CB.Normal($B$5,$B$6)",fill=ASM,align=ctr,font=mono,fmt="0.00%")  # actual return
    C(ws2,f"C{r}",f"=MIN(MAX(B{r},$B$7),$B$8)",align=ctr,fmt="0.00%")                 # credited (clamped)
    C(ws2,f"D{r}",f"=D{r-1}*(1+B{r})",align=ctr,fmt="#,##0")                          # fund grows at actual
    C(ws2,f"E{r}",f"=E{r-1}*(1+C{r})",align=ctr,fmt="#,##0")                          # investor grows at credited
C(ws2,"A23","HOMETOWN PROFIT (Fund − Investor)  →  FORECAST",font=bold,fill=GREY,align=left); ws2.merge_cells("A23:C23")
C(ws2,"D23","=D21-E21",fill=FOR,font=Font(bold=True,size=12),align=ctr,fmt="#,##0")
C(ws2,"A24","Loss? (1 if profit<0)  →  FORECAST (for Q2)",font=bold,fill=GREY,align=left); ws2.merge_cells("A24:C24")
C(ws2,"D24","=IF(D21-E21<0,1,0)",fill=FOR,font=Font(bold=True,size=12),align=ctr)

r=26
C(ws2,f"A{r}","HOW TO USE",font=Font(bold=True,color="1F4E78"),border=False,align=left)
for i,t in enumerate([
 "1. Define B12:B21 (10 annual returns) as Assumptions (CB.Normal 8%,3%).",
 "2. Q1: Define D23 (Hometown profit) as Forecast -> 10,000 trials -> Mean = expected profit; capture frequency chart.",
 "3. Q2: P(Hometown loses) = Certainty of D23 < 0  (set the certainty range to (-inf, 0)).  Or define D24 as a Forecast; its Mean = probability of loss.",
 "Logic: fund grows at the ACTUAL return; investor grows at the CLAMPED rate (min 5%, max 7.5%). Hometown keeps the difference (profit) or covers the shortfall (loss).",
]):
    C(ws2,f"A{r+1+i}",t,border=False,align=left,font=note); ws2.merge_cells(f"A{r+1+i}:F{r+1+i}")
for col,w in {"A":30,"B":24,"C":30,"D":16,"E":18,"F":6}.items(): ws2.column_dimensions[col].width=w
ws2.row_dimensions[10].height=42

out="/home/user/NYU_MSBAi/Simulation_Q4_Q5_Templates.xlsx"
wb.save(out); print("Wrote",out,"sheets:",wb.sheetnames)
